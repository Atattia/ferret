"""Real Arabic OCR/extraction checks; OCR skips if local language data is absent."""
import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from pathlib import Path
import shutil
import tempfile
import unittest

from core.extractor import configure_ocr, extract, extraction_warnings
from core.language import normalize


class ArabicExtractionIntegrationTests(unittest.TestCase):
    def test_utf16_and_docx_table_order(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            plain = root / "arabic.txt"
            plain.write_text("إجازة الموظفين السنوية", encoding="utf-16")
            self.assertEqual(extract(plain), "إجازة الموظفين السنوية")
            document = Document()
            document.add_paragraph("البداية")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "داخل الجدول"
            document.add_paragraph("النهاية")
            path = root / "table.docx"
            document.save(path)
            self.assertEqual(extract(path), "البداية\n\nداخل الجدول\n\nالنهاية")

    @unittest.skipUnless(shutil.which("tesseract") and Path("models/tessdata/ara.traineddata").exists(),
                         "Install Tesseract and run manage.py download-ocr")
    def test_real_scanned_arabic_pdf(self):
        import fitz
        from PyQt6.QtWidgets import QApplication
        from PyQt6.QtGui import QImage, QPainter, QFont
        from PyQt6.QtCore import Qt, QBuffer, QIODevice
        app = QApplication.instance() or QApplication([])
        image = QImage(1800, 300, QImage.Format.Format_RGB32)
        image.fill(Qt.GlobalColor.white)
        painter = QPainter(image)
        painter.setFont(QFont("Arial", 52))
        painter.setPen(Qt.GlobalColor.black)
        painter.drawText(image.rect(), Qt.AlignmentFlag.AlignCenter,
                         "تقرير المدرسة عن إجازة الموظفين ٢٠٢٦")
        painter.end()
        buffer = QBuffer()
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        image.save(buffer, "PNG")
        configure_ocr({"ocr_languages": "ara+eng"})
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "scan.pdf"
            document = fitz.open()
            page = document.new_page(width=720, height=120)
            page.insert_image(page.rect, stream=bytes(buffer.data()))
            document.save(path)
            document.close()
            text = extract(path)
            self.assertIn("المدرسة", normalize(text))
            self.assertIn("الموظفين", normalize(text))
            self.assertEqual(extraction_warnings(), [])


if __name__ == "__main__":
    unittest.main()
